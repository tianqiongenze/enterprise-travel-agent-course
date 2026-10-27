# -*- coding: utf-8 -*-
"""课程文档转换器：全部 md → Word（python-docx）+ PDF（HTML → Edge headless）。

排除：ppt/、tools/、word/、pdf/ 目录本身；代码文件不参与。
用法：python tools/convert.py
"""
import glob
import os
import re
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SKIP_DIRS = {"ppt", "word", "pdf", "tools", ".git"}
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor, Cm  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

PRIMARY = RGBColor(0x06, 0x5A, 0x82)
DARK = RGBColor(0x21, 0x2A, 0x33)


def set_cn_font(run, name="微软雅黑"):
    run.font.name = name
    r = run._element.rPr
    r.rFonts.set(qn("w:eastAsia"), name)


def md_to_docx(md_path: str, docx_path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        # code block
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x33, 0x48, 0x58)
            p.paragraph_format.left_indent = Cm(0.5)
            continue
        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, cell_text in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = re.sub(r"\*\*", "", cell_text)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        set_cn_font(run)
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    cells[j].text = re.sub(r"\*\*(.+?)\*\*", r"\1", val)
                    for para in cells[j].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                            set_cn_font(run)
            doc.add_paragraph()
            continue
        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading("", level=level)
            run = h.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)))
            run.font.color.rgb = PRIMARY if level <= 2 else DARK
            set_cn_font(run, "黑体" if level <= 2 else "微软雅黑")
            i += 1
            continue
        # list items
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
            text = re.sub(r"\[(.)\]", r"\1", text)
            p = doc.add_paragraph(text, style="List Bullet")
            for run in p.runs:
                set_cn_font(run)
            i += 1
            continue
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(1))
            p = doc.add_paragraph(text, style="List Number")
            for run in p.runs:
                set_cn_font(run)
            i += 1
            continue
        if stripped in ("---", "***"):
            doc.add_paragraph()
            i += 1
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            run = p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", stripped.lstrip("> ")))
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x50, 0x60, 0x70)
            set_cn_font(run)
            i += 1
            continue
        if stripped == "":
            i += 1
            continue
        # normal paragraph
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", stripped)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_cn_font(run)
        i += 1

    doc.save(docx_path)


HTML_TMPL = """<!doctype html><html><head><meta charset="utf-8"><style>
body {{ font-family: "Microsoft YaHei", sans-serif; margin: 40px 50px; color: #212A33; line-height: 1.55; }}
h1 {{ color: #065A82; border-bottom: 2px solid #065A82; padding-bottom: 6px; font-size: 22px; }}
h2 {{ color: #065A82; font-size: 18px; margin-top: 24px; }}
h3 {{ color: #1C7293; font-size: 15px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 12px; margin: 8px 0; }}
th {{ background: #E3EEF5; }}
th, td {{ border: 1px solid #C4D3DD; padding: 4px 7px; text-align: left; }}
code, pre {{ font-family: Consolas, monospace; font-size: 11px; background: #F0F4F7; }}
pre {{ padding: 8px 10px; overflow-x: auto; }}
blockquote {{ color: #50606F; border-left: 3px solid #9DB9C8; margin-left: 0; padding-left: 10px; }}
</style></head><body>{body}</body></html>"""


def md_to_pdf(md_path: str, pdf_path: str) -> None:
    body = subprocess.run(
        ["pandoc", md_path, "-t", "html", "--wrap=none"],
        capture_output=True, text=True, encoding="utf-8", check=True,
    ).stdout
    html = HTML_TMPL.format(body=body)
    tmp_html = pdf_path + ".tmp.html"
    open(tmp_html, "w", encoding="utf-8").write(html)
    subprocess.run(
        [EDGE, "--headless", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={pdf_path}", "file:///" + tmp_html.replace("\\", "/")],
        capture_output=True, timeout=120,
    )
    os.remove(tmp_html)


def collect_md_files():
    out = []
    for dirpath, dirnames, filenames in os.walk(ROOT):
        rel = os.path.relpath(dirpath, ROOT)
        top = rel.split(os.sep)[0]
        if top in SKIP_DIRS:
            dirnames[:] = []
            continue
        for fn in filenames:
            if fn.endswith(".md"):
                out.append(os.path.join(dirpath, fn))
    return sorted(out)


def main():
    files = collect_md_files()
    word_dir = os.path.join(ROOT, "word")
    pdf_dir = os.path.join(ROOT, "pdf")
    os.makedirs(word_dir, exist_ok=True)
    os.makedirs(pdf_dir, exist_ok=True)
    ok = fail = 0
    for md_path in files:
        rel = os.path.relpath(md_path, ROOT)
        safe = re.sub(r"[^\w.\-一-龥]+", "_", rel[:-3])
        docx_path = os.path.join(word_dir, safe + ".docx")
        pdf_path = os.path.join(pdf_dir, safe + ".pdf")
        try:
            md_to_docx(md_path, docx_path)
            ok += 1
        except Exception as e:  # noqa: BLE001
            print("DOCX FAIL:", rel, e)
            fail += 1
        try:
            md_to_pdf(md_path, pdf_path)
            ok += 1
        except Exception as e:  # noqa: BLE001
            print("PDF  FAIL:", rel, e)
            fail += 1
    print(f"converted: ok={ok} fail={fail} (files={len(files)})")


if __name__ == "__main__":
    main()
